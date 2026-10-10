"""数据入库：教材 PDF -> 结构化分块 -> 本地 bge-m3 嵌入 -> ChromaDB 持久化。

用法：
    python -m backend.rag.ingest             # 入库默认课程语料
    python -m backend.rag.ingest --reset     # 先清空 collection 再重建
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from backend.core.config import settings
from backend.rag.chunking import (Chunk, chunk_by_heading, chunk_fixed_overlap,
                                  chunk_pages, parse_structured_pdf)

# chromadb / llama_index / torch 一律**延迟导入**：本模块的解析端（discover_corpus / parse_one）
# 只依赖 PyMuPDF + python-pptx，被 scripts/build_fts_from_raw.py 复用来建 SQLite FTS5 索引。
# 早先这几个 import 写在模块顶层，导致「不装向量栈就 import 不了解析函数」，
# 逼着重复实现一份解析逻辑。

# 课程语料。parser: structured=按章节切文本 PDF；fixed=固定长度切排版乱文档；pptx=课件
# 遗留的数据结构三本（扁平放在 data/raw/ 下），course=ds-algo
DEFAULT_CORPUS = [
    {"file": "hello-algo_zh_c.pdf", "name": "Hello 算法",
     "doc_type": "textbook", "parser": "structured", "course": "ds-algo"},
    {"file": "open-data-structures.pdf", "name": "Open Data Structures",
     "doc_type": "textbook", "parser": "structured", "course": "ds-algo"},
    {"file": "cp-handbook.pdf", "name": "Competitive Programmer's Handbook",
     "doc_type": "exercise", "parser": "structured", "course": "ds-algo"},
]

# 目录约定（见 docs/07-408数据采集清单）：data/raw/<课程目录>/<类别>/*.pdf|pptx
# 这四个是历史遗留的短目录名；课程搬进数据库之后，新课一律用 course_id 本身做目录名
# （见 _course_dirs），否则教务新开的课传了资料也扫不到。
DIR2COURSE = {"ds": "ds-algo", "co": "co", "os": "os", "net": "net"}
# 进 RAG 的类别（past_papers 走题库入库，不进检索语料）→ parser
RAG_DOCTYPE_PARSER = {"textbook": "structured", "exercise": "structured", "courseware": "fixed"}
# 后缀 → parser 的覆盖项。空值表示"用类别决定"（PDF 就是这一档：
# 教材按章节切、课件按页切）。上传接口的 _EXT_PARSER 是同一份名单的另一面，
# 由 test_material_ingest 钉住两者一致 —— 这里少一个后缀，整轮重建就会
# 静默丢掉教务传上来的那类文件，而且列表里看不出少了什么。
EXT_PARSER: dict[str, str] = {".pdf": "", ".pptx": "pptx", ".ppt": "pptx",
                              ".docx": "docx", ".txt": "text"}


def _course_dirs() -> dict[str, str]:
    """目录名 → course_id。历史短名 + 当前课表里每门课的 id 本身。

    读课表失败时（库还没建、纯解析场景）退回历史短名 —— 这个函数服务的是
    离线解析，不该因为连不上库就整个扫不出东西。
    """
    dirs = dict(DIR2COURSE)
    try:
        from backend.core.courses import course_ids
        for cid in course_ids():
            dirs.setdefault(cid, cid)
    except Exception:  # noqa: BLE001
        pass
    return dirs


def discover_corpus(raw_dir: Path) -> list[dict]:
    """扫描 data/raw/<课程目录>/<textbook|exercise|courseware>/ 下的 pdf/pptx，自动生成语料清单。"""
    items: list[dict] = []
    for d, cid in _course_dirs().items():
        for sub, parser in RAG_DOCTYPE_PARSER.items():
            base = raw_dir / d / sub
            if not base.is_dir():
                continue
            for f in sorted(base.iterdir()):
                ext = f.suffix.lower()
                if ext not in EXT_PARSER:
                    continue
                p = EXT_PARSER[ext] or parser
                items.append({"file": str(f.relative_to(raw_dir)), "name": f.stem,
                              "doc_type": sub, "parser": p, "course": cid})
    return items


def _pptx_slides(path: str) -> list[str]:
    """抽取 PPTX 每页文本，按幻灯片顺序返回。未装 python-pptx 时提示并返回空列表。"""
    try:
        from pptx import Presentation
    except ImportError:
        print("  [跳过 pptx] 需要 python-pptx：pip install python-pptx")
        return []
    prs = Presentation(path)
    slides = []
    for slide in prs.slides:
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append("\n".join(p.text for p in shape.text_frame.paragraphs if p.text))
        slides.append("\n".join(x for x in parts if x.strip()))
    return slides


# Word 的标题样式名：英文版是 "Heading 1"，中文版是 "标题 1"，
# 还有些模板留着小写的 "heading 1"。三种都认，认不出就当正文。
_RE_DOCX_HEADING = re.compile(r"^(?:heading|标题)\s*(\d+)", re.I)
# 纯文本没有样式，只能按行猜标题。两个最稳的形状：中文教材的「第3章/第三节」，
# 和数字编号的「3.2 xxx」。猜不中的后果只是少一层定位，不会切错内容。
_RE_TXT_CHAPTER = re.compile(r"^\s*第\s*[0-9〇一二三四五六七八九十百]+\s*[章讲部编]\s*\S")
_RE_TXT_SECTION = re.compile(r"^\s*(?:\d+\.\d+(?:\.\d+)?|第\s*[0-9〇一二三四五六七八九十百]+\s*[节小])\s*\S")
_TXT_HEADING_MAX = 60      # 再长就不像标题了，多半是正文里恰好以编号开头的一句


def _docx_blocks(path: str) -> list[tuple[int, str]]:
    """抽取 .docx 的 (标题级别, 文本)，正文级别为 0。未装 python-docx 时提示并返回空。

    按 body 的 XML 子节点顺序遍历，而不是先 doc.paragraphs 再 doc.tables ——
    后者会把全文的表格都堆到末尾，习题类文档（题干在段落、选项在表格）就此错位。

    老的二进制 .doc 不在此列：python-docx 读不了它，得另存为 .docx。
    """
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError:
        print("  [跳过 docx] 需要 python-docx：pip install python-docx")
        return []

    doc = Document(path)
    blocks: list[tuple[int, str]] = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style else "") or ""
            m = _RE_DOCX_HEADING.match(style)
            blocks.append((int(m.group(1)) if m else 0, text))
        elif child.tag == qn("w:tbl"):
            for row in Table(child, doc).rows:
                cells = [c.text.strip() for c in row.cells]
                line = " | ".join(x for x in cells if x)
                if line:
                    blocks.append((0, line))
    return blocks


def _text_blocks(path: str) -> list[tuple[int, str]]:
    """读纯文本并按行猜标题层级。

    编码依次试 utf-8 → gbk → utf-8(replace)。GBK 这一档不能省：国内老师手上的
    txt 有相当一部分是 GBK，直接按 utf-8 解会整篇变成乱码 —— 而乱码不会报错，
    它会安安静静地进索引，然后永远检不到。
    """
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")

    blocks: list[tuple[int, str]] = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        level = 0
        if len(stripped) <= _TXT_HEADING_MAX:
            if _RE_TXT_CHAPTER.match(stripped):
                level = 1
            elif _RE_TXT_SECTION.match(stripped):
                level = 2
        blocks.append((level, stripped))
    return blocks


def build_embed_model():
    from llama_index.embeddings.huggingface import HuggingFaceEmbedding

    from backend.core.device import load_with_fallback

    model, _ = load_with_fallback(
        lambda d: HuggingFaceEmbedding(model_name=settings.embed_model_path, device=d,
                                       trust_remote_code=True),
        settings.device, label="嵌入模型 bge-m3")
    return model


def get_chroma_collection(reset: bool = False):
    import chromadb

    Path(settings.chroma_dir).mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=settings.chroma_dir)
    if reset:
        try:
            client.delete_collection(settings.chroma_collection)
            print(f"[reset] 已删除旧 collection: {settings.chroma_collection}")
        except Exception:
            pass
    return client.get_or_create_collection(settings.chroma_collection)


def chunks_to_nodes(chunks: list[Chunk]) -> list:
    from llama_index.core.schema import TextNode

    nodes = []
    skip = ["page_start", "page_end", "chapter", "section", "doc_type", "chunk_index", "course"]
    for c in chunks:
        node = TextNode(text=c.text, metadata=c.to_metadata())
        node.excluded_embed_metadata_keys = skip
        node.excluded_llm_metadata_keys = skip
        nodes.append(node)
    return nodes


def parse_one(item: dict, raw_dir: Path) -> list[Chunk]:
    path = raw_dir / item["file"]
    if not path.exists():
        print(f"[跳过] 未找到文件: {path}")
        return []
    name, doc_type, parser = item["name"], item["doc_type"], item["parser"]
    course = item.get("course", "")
    print(f"[解析] {name} <- {item['file']}  (parser={parser}, course={course})")
    if parser == "structured":
        chunks = parse_structured_pdf(str(path), name, doc_type=doc_type, course=course)
    elif parser == "fixed":
        import fitz
        # 按页聚合而不是拼成一根长字符串：课件的引用必须能指回具体页码
        pages = [p.get_text() for p in fitz.open(str(path))]
        chunks = chunk_pages(pages, name, doc_type=doc_type, course=course)
        if not chunks:                       # 整份没有文本层（扫描件）时退回原路径，至少不丢内容
            chunks = chunk_fixed_overlap("\n".join(pages), name, doc_type=doc_type, course=course)
    elif parser == "pptx":
        slides = _pptx_slides(str(path))
        chunks = chunk_pages(slides, name, doc_type=doc_type, course=course,
                             label="第 {a}~{b} 张幻灯片")
    elif parser in ("docx", "text"):
        # 这两类没有页码，引用里不会出现"第 87 页"，标题路径是唯一的定位物
        blocks = _docx_blocks(str(path)) if parser == "docx" else _text_blocks(str(path))
        chunks = chunk_by_heading(blocks, name, doc_type=doc_type, course=course)
    else:
        raise ValueError(f"未知 parser 类型: {parser}")
    print(f"        得到 {len(chunks)} 个分块")
    return chunks


def ingest(corpus: list[dict] | None = None, reset: bool = False,
           course: str | None = None) -> int:
    raw_dir = Path(settings.raw_dir)
    if corpus is None:
        # 遗留扁平 ds 语料 + 自动发现的 data/raw/<课程>/<类别>/ 语料
        corpus = DEFAULT_CORPUS + discover_corpus(raw_dir)
    if course:
        corpus = [it for it in corpus if it.get("course") == course]
    all_chunks: list[Chunk] = []
    for item in corpus:
        all_chunks.extend(parse_one(item, raw_dir))
    if not all_chunks:
        print("没有可入库的内容。")
        return 0

    from llama_index.core import StorageContext, VectorStoreIndex
    from llama_index.vector_stores.chroma import ChromaVectorStore

    print(f"[嵌入] 共 {len(all_chunks)} 个分块，加载本地 bge-m3 ...")
    embed_model = build_embed_model()
    nodes = chunks_to_nodes(all_chunks)

    collection = get_chroma_collection(reset=reset)
    vector_store = ChromaVectorStore(chroma_collection=collection)
    storage_context = StorageContext.from_defaults(vector_store=vector_store)

    print("[入库] 写入 ChromaDB ...")
    VectorStoreIndex(nodes, storage_context=storage_context, embed_model=embed_model, show_progress=True)
    print(f"[完成] 已写入 {len(nodes)} 个节点到 '{settings.chroma_collection}' @ {settings.chroma_dir}")
    return len(nodes)


def main():
    ap = argparse.ArgumentParser(description="CourseLoop RAG 数据入库（含 408 四门自动发现）")
    ap.add_argument("--reset", action="store_true", help="入库前清空已有 collection")
    ap.add_argument("--course", default=None, help="只入某门：ds-algo|co|os|net")
    args = ap.parse_args()
    ingest(reset=args.reset, course=args.course)


if __name__ == "__main__":
    main()
